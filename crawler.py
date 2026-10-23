"""
공지·공고·대학원 입시 크롤러 (핵심 모듈)

두 가지 방식으로 동작한다.
  1) board : 게시판 목록을 읽어 새 글만 골라낸다.
  2) watch : 게시판이 없는 안내 페이지를 통째로 저장해 두고, 바뀌면 새로 생긴 문장을 알린다.

결과는 notices.db (SQLite) 에 저장된다. 웹 대시보드(app.py)가 이 파일을 읽는다.

명령줄 실행:
  python crawler.py                 전체 확인
  python crawler.py --group 일본    분류 하나만
  python crawler.py --export        notices.xlsx 로 내보내기
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sqlite3
import sys
import threading
import time
from datetime import datetime, timedelta
from pathlib import Path
from typing import Callable
from urllib.parse import urljoin, urlparse
from urllib.robotparser import RobotFileParser

import requests
import yaml
from bs4 import BeautifulSoup

BASE_DIR = Path(__file__).resolve().parent
DB_PATH = BASE_DIR / "notices.db"
CONFIG_PATH = BASE_DIR / "config.yaml"
SNAP_DIR = BASE_DIR / "page_snapshots"
XLSX_PATH = BASE_DIR / "notices.xlsx"

# 신원을 숨기지 않는다: 크롤러임을 밝히는 User-Agent
USER_AGENT = "PersonalNoticeCrawler/2.0 (personal use; low frequency)"
# 2026-10-23 / 2026.10.23 / 2026. 10. 23. / 2026/10/23 / 2026年10月23日 모두 인식
DATE_RE = re.compile(r"(20\d{2})\s*[.\-/年]\s*(\d{1,2})\s*[.\-/月]\s*(\d{1,2})")


def now_str() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def clean(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def norm_date(text: str) -> str:
    m = DATE_RE.search(text or "")
    if not m:
        return ""
    y, mo, d = m.groups()
    return f"{y}-{int(mo):02d}-{int(d):02d}"


def load_config(path: Path = CONFIG_PATH) -> dict:
    return yaml.safe_load(path.read_text(encoding="utf-8"))


# ================================================================ 설정 편집 (대시보드의 관리 화면용)
CONFIG_HEADER = """# 공지·입시 감시 설정 파일
# 대시보드의 '감시 대상 관리' 화면에서 추가·수정·삭제하면 이 파일이 자동으로 다시 쓰인다.
# 직접 고쳐도 되지만, 화면에서 저장하면 이 파일의 주석은 이 머리말만 남는다. 설명은 각 대상의 note 칸에 적는다.
# 저장할 때마다 바로 이전 내용이 config.yaml.bak 에 남는다.
"""
FIELD_ORDER = ["name", "group", "sub", "type", "enabled", "note", "list_url", "url", "pages", "page_size",
               "row_selector", "title_selector", "title_text_selector", "date_selector",
               "remove_in_title", "category_selector", "categories", "link_regex", "link_template",
               "detail_selector", "content_selector", "remove_selectors", "ignore_regex", "keywords",
               "include", "exclude", "require", "render", "wait_selector"]
LIST_FIELDS = {"categories", "keywords", "include", "exclude", "require", "remove_selectors"}
_config_lock = threading.Lock()


def save_config(cfg: dict, path: Path = CONFIG_PATH) -> None:
    """바로 이전 파일을 .bak 로 남기고, 임시 파일에 쓴 뒤 바꿔 끼운다 (쓰다 꺼져도 설정이 깨지지 않게)."""
    text = CONFIG_HEADER + yaml.safe_dump(cfg, allow_unicode=True, sort_keys=False, width=1000)
    yaml.safe_load(text)  # 다시 읽히는지 확인
    with _config_lock:
        if path.exists():
            path.with_suffix(".yaml.bak").write_text(path.read_text(encoding="utf-8"), encoding="utf-8")
        tmp = path.with_suffix(".yaml.tmp")
        tmp.write_text(text, encoding="utf-8")
        tmp.replace(path)


def clean_source(raw: dict) -> dict:
    """화면에서 받은 값을 정리하고 검사한다. 문제가 있으면 ValueError(한국어 설명)."""
    src: dict = {}
    for k, v in raw.items():
        if k not in FIELD_ORDER:
            continue
        if k in LIST_FIELDS:
            if isinstance(v, str):
                v = [x.strip() for x in re.split(r"[,\n]", v) if x.strip()]
            v = [str(x).strip() for x in (v or []) if str(x).strip()]
        elif k == "pages":
            try:
                v = max(1, min(10, int(v or 1)))
            except (TypeError, ValueError):
                raise ValueError("페이지 수는 1~10 사이 숫자여야 한다")
        elif k == "enabled":
            v = bool(v)
        elif isinstance(v, str):
            v = v.strip()
        if v in ("", None, []):
            continue
        src[k] = v
    src.setdefault("type", "board")
    src.setdefault("enabled", True)
    if not src.get("name"):
        raise ValueError("이름을 입력해야 한다")
    if not src.get("group"):
        raise ValueError("분류를 입력해야 한다 (예: 일본)")
    if src["type"] not in ("board", "watch"):
        raise ValueError("방식은 board 또는 watch 여야 한다")
    url_key = "list_url" if src["type"] == "board" else "url"
    other = "url" if url_key == "list_url" else "list_url"
    if not src.get(url_key) and src.get(other):  # 방식을 바꿨을 때 주소를 옮겨 준다
        src[url_key] = src.pop(other)
    src.pop(other, None)
    u = src.get(url_key, "")
    if not re.match(r"^https?://[^\s/]+\.[^\s]+$", u):
        raise ValueError("주소는 http:// 또는 https:// 로 시작하는 전체 주소여야 한다")
    if src["type"] == "board":
        src.setdefault("row_selector", "table tbody tr")
        src.setdefault("title_selector", "a")
        src.setdefault("pages", 1)
        if src["pages"] > 1 and "{page}" not in u and "{offset}" not in u:
            raise ValueError("2페이지 이상 보려면 주소의 페이지 번호 자리에 {page} (또는 글 순번 자리에 {offset}) 를 넣어야 한다")
        if "page_size" in src:
            try:
                src["page_size"] = max(1, min(200, int(src["page_size"])))
            except (TypeError, ValueError):
                raise ValueError("page_size 는 숫자여야 한다")
        for k in ("content_selector", "remove_selectors", "ignore_regex", "keywords"):
            src.pop(k, None)
    else:
        for k in ("pages", "page_size", "row_selector", "title_selector", "title_text_selector", "date_selector",
                  "remove_in_title", "category_selector", "categories", "link_regex",
                  "link_template", "include", "exclude", "require", "detail_selector"):
            src.pop(k, None)
    for k in ("link_regex", "ignore_regex"):
        if src.get(k):
            try:
                re.compile(src[k])
            except re.error as e:
                raise ValueError(f"{k} 정규식 오류: {e}")
    sel_keys = ("row_selector", "title_selector", "title_text_selector", "date_selector",
                "remove_in_title", "category_selector", "content_selector", "detail_selector")
    probe = BeautifulSoup("<div></div>", "html.parser")
    for k in sel_keys:
        if src.get(k):
            try:
                probe.select(src[k])
            except Exception:
                raise ValueError(f"{k} 의 CSS 선택자 형식이 잘못됐다: {src[k]}")
    return {k: src[k] for k in FIELD_ORDER if k in src}


KEEP_ON_EDIT = ("page_size", "render", "wait_selector")   # 화면 폼에 없는 값: 수정해도 지우지 않는다


def upsert_source(raw: dict, old_name: str | None) -> dict:
    cfg = load_config()
    old = next((s for s in cfg["sources"] if s["name"] == old_name), None) if old_name else None
    if old:
        raw = {**{k: old[k] for k in KEEP_ON_EDIT if k in old and k not in raw}, **raw}
    src = clean_source(raw)
    names = [s["name"] for s in cfg["sources"]]
    if src["name"] != old_name and src["name"] in names:
        raise ValueError(f"'{src['name']}' 이름이 이미 있다. 다른 이름을 쓸 것")
    if old_name and old_name in names:
        cfg["sources"][names.index(old_name)] = src
    else:
        cfg["sources"].append(src)
    save_config(cfg)
    return src


def delete_source(name: str) -> None:
    cfg = load_config()
    before = len(cfg["sources"])
    cfg["sources"] = [s for s in cfg["sources"] if s["name"] != name]
    if len(cfg["sources"]) == before:
        raise ValueError(f"'{name}' 대상을 찾지 못했다")
    save_config(cfg)


def set_enabled(name: str, enabled: bool) -> None:
    cfg = load_config()
    for s in cfg["sources"]:
        if s["name"] == name:
            s["enabled"] = bool(enabled)
            save_config(cfg)
            return
    raise ValueError(f"'{name}' 대상을 찾지 못했다")


def move_source(name: str, delta: int) -> None:
    cfg = load_config()
    names = [s["name"] for s in cfg["sources"]]
    i = names.index(name)
    j = max(0, min(len(names) - 1, i + delta))
    cfg["sources"].insert(j, cfg["sources"].pop(i))
    save_config(cfg)


# ================================================================ 저장소
class Store:
    """SQLite 저장소. 웹서버의 여러 스레드에서 같이 쓰므로 잠금으로 보호한다."""

    def __init__(self, path: Path = DB_PATH):
        self.lock = threading.Lock()
        self.db = sqlite3.connect(path, check_same_thread=False)
        self.db.row_factory = sqlite3.Row
        self.db.executescript("""
        CREATE TABLE IF NOT EXISTS items (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            key TEXT UNIQUE, tkey TEXT,
            grp TEXT, source TEXT, date TEXT, title TEXT, url TEXT, change TEXT,
            kind TEXT, found_at TEXT, is_read INTEGER DEFAULT 0, starred INTEGER DEFAULT 0);
        CREATE INDEX IF NOT EXISTS idx_tkey ON items(tkey);
        CREATE TABLE IF NOT EXISTS sources (
            name TEXT PRIMARY KEY, grp TEXT, kind TEXT, url TEXT,
            last_checked TEXT, last_ok TEXT, last_error TEXT, last_new INTEGER DEFAULT 0);
        CREATE TABLE IF NOT EXISTS meta (k TEXT PRIMARY KEY, v TEXT);
        """)
        # 공지 정리·번역 칸 (예전 DB 에는 없으므로 없으면 추가)
        have = {r["name"] for r in self.db.execute("PRAGMA table_info(items)")}
        for col, typ in [("summary", "TEXT"), ("title_ko", "TEXT"), ("sum_mode", "TEXT"),
                         ("sum_state", "TEXT"), ("sum_tries", "INTEGER DEFAULT 0"),
                         ("sum_error", "TEXT"), ("sum_at", "TEXT")]:
            if col not in have:
                self.db.execute(f"ALTER TABLE items ADD COLUMN {col} {typ}")
        self.db.commit()

    def has_key(self, key: str) -> bool:
        with self.lock:
            return self.db.execute("SELECT 1 FROM items WHERE key=?", (key,)).fetchone() is not None

    # ---- 공지 정리
    def pending_summaries(self, mode: str, max_tries: int) -> list[dict]:
        """정리할 글. 키 없이 발췌만 해 둔 글은 키가 생기면(mode=llm) 다시 정리한다."""
        sql = ("SELECT * FROM items WHERE (sum_mode IS NULL AND COALESCE(sum_tries,0) < ?)"
               + (" OR sum_mode='extract'" if mode == "llm" else "")
               + " ORDER BY COALESCE(NULLIF(date,''), substr(found_at,1,10)) DESC, id DESC")
        with self.lock:
            return [dict(r) for r in self.db.execute(sql, (max_tries,)).fetchall()]

    def set_summary(self, item_id: int, mode: str, title_ko: str, points: list[str], state: str) -> None:
        with self.lock:
            self.db.execute(
                "UPDATE items SET summary=?, title_ko=?, sum_mode=?, sum_state=?, sum_error=NULL, sum_at=?"
                " WHERE id=?",
                (json.dumps(points, ensure_ascii=False), title_ko or None, mode, state, now_str(), item_id))
            self.db.commit()

    def summary_failed(self, item_id: int, error: str) -> None:
        with self.lock:
            self.db.execute("UPDATE items SET sum_tries=COALESCE(sum_tries,0)+1, sum_error=? WHERE id=?",
                            (error, item_id))
            self.db.commit()

    # ---- 크롤 결과 기록
    def source_seen_before(self, name: str) -> bool:
        with self.lock:
            r = self.db.execute("SELECT last_ok FROM sources WHERE name=?", (name,)).fetchone()
        return bool(r and r["last_ok"])

    def add_item(self, it: dict, mark_read: bool) -> bool:
        """새 항목이면 넣고 True. 같은 링크 또는 같은 제목+날짜가 이미 있으면 False."""
        with self.lock:
            dup = self.db.execute(
                "SELECT 1 FROM items WHERE key=? OR (tkey IS NOT NULL AND tkey=?)",
                (it["key"], it.get("tkey")),
            ).fetchone()
            if dup:
                return False
            self.db.execute(
                "INSERT INTO items(key,tkey,grp,source,date,title,url,change,kind,found_at,is_read)"
                " VALUES(?,?,?,?,?,?,?,?,?,?,?)",
                (it["key"], it.get("tkey"), it["group"], it["source"], it["date"], it["title"],
                 it["url"], it.get("change", ""), it["kind"], now_str(), 1 if mark_read else 0),
            )
            self.db.commit()
            return True

    def set_source(self, src: dict, ok: bool, error: str = "", new: int = 0) -> None:
        with self.lock:
            prev = self.db.execute("SELECT last_ok FROM sources WHERE name=?", (src["name"],)).fetchone()
            last_ok = now_str() if ok else (prev["last_ok"] if prev else None)
            self.db.execute(
                "INSERT INTO sources(name,grp,kind,url,last_checked,last_ok,last_error,last_new)"
                " VALUES(?,?,?,?,?,?,?,?) ON CONFLICT(name) DO UPDATE SET grp=excluded.grp,"
                " kind=excluded.kind, url=excluded.url, last_checked=excluded.last_checked,"
                " last_ok=excluded.last_ok, last_error=excluded.last_error, last_new=excluded.last_new",
                (src["name"], src.get("group", ""), src.get("type", "board"),
                 src.get("url") or src.get("list_url"), now_str(), last_ok, error, new),
            )
            self.db.commit()

    def set_meta(self, k: str, v: str) -> None:
        with self.lock:
            self.db.execute("INSERT OR REPLACE INTO meta(k,v) VALUES(?,?)", (k, v))
            self.db.commit()

    def get_meta(self, k: str) -> str | None:
        with self.lock:
            r = self.db.execute("SELECT v FROM meta WHERE k=?", (k,)).fetchone()
        return r["v"] if r else None

    # ---- 대시보드 조회
    def items(self, group: str = "", q: str = "", unread: bool = False,
              starred: bool = False, limit: int = 400) -> list[dict]:
        sql, args = "SELECT * FROM items WHERE 1=1", []
        if group:
            sql += " AND grp=?"; args.append(group)
        if q:
            sql += (" AND (title LIKE ? OR source LIKE ? OR change LIKE ?"
                    " OR COALESCE(title_ko,'') LIKE ? OR COALESCE(summary,'') LIKE ?)")
            args += [f"%{q}%"] * 5
        if unread:
            sql += " AND is_read=0"
        if starred:
            sql += " AND starred=1"
        sql += " ORDER BY is_read ASC, COALESCE(NULLIF(date,''), substr(found_at,1,10)) DESC, id DESC LIMIT ?"
        args.append(limit)
        with self.lock:
            return [dict(r) for r in self.db.execute(sql, args).fetchall()]

    def counts(self) -> dict:
        with self.lock:
            rows = self.db.execute(
                "SELECT grp, COUNT(*) n, SUM(is_read=0) unread FROM items GROUP BY grp").fetchall()
        return {r["grp"]: {"total": r["n"], "unread": r["unread"] or 0} for r in rows}

    def sources(self) -> list[dict]:
        with self.lock:
            return [dict(r) for r in self.db.execute("SELECT * FROM sources").fetchall()]

    def mark(self, ids: list[int], field: str, value: int) -> None:
        assert field in ("is_read", "starred")
        with self.lock:
            self.db.executemany(f"UPDATE items SET {field}=? WHERE id=?", [(value, i) for i in ids])
            self.db.commit()

    def forget_source(self, name: str, purge_items: bool) -> None:
        """대상을 지울 때: 상태 기록은 지우고, 원하면 그동안 모은 글도 지운다."""
        with self.lock:
            self.db.execute("DELETE FROM sources WHERE name=?", (name,))
            if purge_items:
                self.db.execute("DELETE FROM items WHERE source=?", (name,))
            self.db.commit()
        p = snap_path(name)
        if p.exists():
            p.unlink()

    def rename_source(self, old: str, new: str, group: str) -> None:
        with self.lock:
            self.db.execute("UPDATE items SET source=?, grp=? WHERE source=?", (new, group, old))
            self.db.execute("UPDATE sources SET name=?, grp=? WHERE name=?", (new, group, old))
            self.db.commit()
        p_old, p_new = snap_path(old), snap_path(new)
        if p_old.exists() and not p_new.exists():
            p_old.rename(p_new)

    def mark_all_read(self, group: str = "") -> None:
        with self.lock:
            if group:
                self.db.execute("UPDATE items SET is_read=1 WHERE grp=?", (group,))
            else:
                self.db.execute("UPDATE items SET is_read=1")
            self.db.commit()


# ================================================================ robots.txt
_robots_cache: dict[str, RobotFileParser | None] = {}


def robots_allows(url: str, session: requests.Session) -> bool:
    """robots.txt 가 명시적으로 막으면 False. 파일이 없거나 HTML 오류 페이지면 허용."""
    p = urlparse(url)
    host = f"{p.scheme}://{p.netloc}"
    if host not in _robots_cache:
        rp = None
        try:
            r = session.get(host + "/robots.txt", timeout=10)
            body = r.text
            if r.status_code == 200 and "<html" not in body.lower()[:500]:
                rp = RobotFileParser()
                rp.parse(body.splitlines())
        except requests.RequestException:
            pass
        _robots_cache[host] = rp
    rp = _robots_cache[host]
    return True if rp is None else rp.can_fetch(USER_AGENT, url)


# ================================================================ 가져오기
def fetch_static(url: str, session: requests.Session) -> str:
    r = session.get(url, timeout=25)
    r.raise_for_status()
    if not r.encoding or r.encoding.lower() == "iso-8859-1":
        r.encoding = r.apparent_encoding
    return r.text


def fetch_rendered(url: str, wait_selector: str | None) -> str:
    """JavaScript 로 내용을 그리는 사이트용. (pip install playwright; playwright install chromium)"""
    from playwright.sync_api import sync_playwright

    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(user_agent=USER_AGENT)
        page.goto(url, wait_until="networkidle", timeout=60_000)
        if wait_selector:
            page.wait_for_selector(wait_selector, timeout=30_000)
        html = page.content()
        browser.close()
    return html


def fetch(url: str, src: dict, session: requests.Session) -> str:
    if src.get("render") == "browser":
        return fetch_rendered(url, src.get("wait_selector"))
    return fetch_static(url, session)


# ================================================================ board
def build_link(raw_href: str, page_url: str, src: dict) -> str:
    rx, tpl = src.get("link_regex"), src.get("link_template")
    if rx and tpl:
        m = re.search(rx, raw_href or "")
        if m:
            return tpl.format(*m.groups())
    if raw_href and not raw_href.lower().startswith("javascript"):
        return urljoin(page_url, raw_href)
    return page_url


def parse_list(html: str, page_url: str, src: dict) -> list[dict]:
    soup = BeautifulSoup(html, "html.parser")
    cats = src.get("categories") or []
    items = []
    for row in soup.select(src.get("row_selector", "table tbody tr")):
        a = row.select_one(src.get("title_selector", "a"))
        if a is None:
            continue
        if cats and src.get("category_selector"):
            cat_el = row.select_one(src["category_selector"])
            cat = clean(cat_el.get_text(" ")) if cat_el else ""
            if not any(c in cat for c in cats):
                continue
        date_el = row.select_one(src["date_selector"]) if src.get("date_selector") else None
        date = norm_date(date_el.get_text(" ") if date_el else row.get_text(" "))
        href = a.get("href", "")
        t_el = a.select_one(src["title_text_selector"]) if src.get("title_text_selector") else a
        t_el = t_el or a
        if src.get("remove_in_title"):
            for tag in t_el.select(src["remove_in_title"]):
                tag.decompose()
        title = clean(t_el.get_text(" "))
        if title:
            items.append({"title": title, "date": date, "url": build_link(href, page_url, src)})
    return items


def keyword_ok(title: str, src: dict) -> bool:
    """include: 이 중 하나라도 있으면 통과 / require: 그와 별도로 이 중 하나도 있어야 통과 / exclude: 하나라도 있으면 숨김"""
    inc, exc, req = src.get("include") or [], src.get("exclude") or [], src.get("require") or []
    t = title.lower()
    if inc and not any(k.lower() in t for k in inc):
        return False
    if req and not any(k.lower() in t for k in req):
        return False
    return not any(k.lower() in t for k in exc)


def topic_hits(it: dict, topics: list[dict]) -> list[str]:
    """글 제목이 걸리는 주제(채용·상담 등) 이름들. 화면의 '주제 켜기/끄기'에 쓴다."""
    t = (str(it.get("title", "")) + " " + str(it.get("title_ko") or "")).lower()
    grp = it.get("grp") or it.get("group")
    return [tp["name"] for tp in topics or []
            if (not tp.get("groups") or grp in tp["groups"])
            and any(str(k).lower() in t for k in tp.get("keywords") or [])]


def visible_filter(cfg: dict, hidden_topics: set[str] | None = None):
    """현재 설정의 include/exclude/require 를 이미 모은 게시판 글에도 적용하는 판별 함수.
    필터를 나중에 바꿔도 DB 는 지우지 않고 화면에서만 가린다 (필터를 풀면 다시 보인다).
    hidden_topics 에 든 주제에 걸리는 글도 가린다."""
    by_name = {s.get("name"): s for s in cfg.get("sources", [])}
    topics = cfg.get("topics") or []
    hidden = set(hidden_topics or ())
    cutoff = cutoff_date(cfg)

    def ok(it: dict) -> bool:
        if cutoff and it.get("date") and it["date"] < cutoff:   # max_age_months 보다 오래된 글
            return False
        if hidden and hidden & set(topic_hits(it, topics)):
            return False
        src = by_name.get(it.get("source"))
        if not src or it.get("kind") != "board":
            return True
        return keyword_ok(it.get("title", ""), src)
    return ok


def page_url(src: dict, i: int) -> str:
    """i 번째 목록 페이지 주소. {page}=페이지 번호, {offset}=(i-1)*page_size (글 순번으로 넘기는 사이트용)."""
    return src["list_url"].replace("{page}", str(i)).replace("{offset}", str((i - 1) * int(src.get("page_size", 10))))


def cutoff_date(cfg: dict) -> str:
    """max_age_months(기본 8)개월 전 날짜 'YYYY-MM-DD'. 0 이면 '' (제한 없음)."""
    months = int(cfg.get("max_age_months", 8) or 0)
    if months <= 0:
        return ""
    t = datetime.now()
    y, m = divmod(t.month - 1 - months, 12)
    y, m = t.year + y, m + 1
    last_day = [31, 29 if y % 4 == 0 and (y % 100 or y % 400 == 0) else 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31][m - 1]
    return f"{y}-{m:02d}-{min(t.day, last_day):02d}"


def crawl_board(src: dict, session: requests.Session, delay: float,
                cutoff: str = "", max_pages: int = 25,
                is_known: Callable[[str], bool] | None = None) -> list[dict]:
    """게시판 목록을 읽는다. cutoff 가 있으면 그 날짜까지 페이지를 더 넘겨 빠짐없이 모으고,
    그보다 오래된 글은 버린다. 페이지 수는 max_pages 를 넘지 않는다 (서버 부담 제한)."""
    found, seen = [], set()
    min_pages = int(src.get("pages", 1))
    can_page = "{page}" in src["list_url"] or "{offset}" in src["list_url"]
    limit = max(min_pages, max_pages) if (cutoff and can_page) else min_pages
    prev_keys: set = set()
    earlier: set = set()   # 앞 페이지에서 본 글 (상단 고정글은 매 페이지 반복된다)
    for i in range(1, limit + 1):
        url = page_url(src, i)
        if not robots_allows(url, session):
            raise PermissionError(f"robots.txt 가 수집을 막고 있다: {url}")
        page_items = parse_list(fetch(url, src, session), url, src)
        time.sleep(delay)
        if not page_items:
            break
        keys = {it["url"] if it["url"] != url else it["title"] for it in page_items}
        if keys == prev_keys:  # 페이지 번호를 무시하고 같은 목록을 주는 사이트
            break
        prev_keys = keys
        for it in page_items:
            if cutoff and it["date"] and it["date"] < cutoff:
                continue
            if not keyword_ok(it["title"], src):
                continue
            key = it["url"] if it["url"] != url else it["title"]
            if key in seen:  # 상단 고정글이 페이지마다 반복
                continue
            seen.add(key)
            found.append({**it, "key": key, "tkey": f"T|{it['title']}|{it['date']}",
                          "change": "", "kind": "board"})
        fresh = [it for it in page_items if (it["url"] if it["url"] != url else it["title"]) not in earlier]
        earlier |= keys
        # 고정글은 첫 페이지에서 이미 봤으므로, 이 페이지에 처음 나온 글의 날짜로 판단한다
        dates = [it["date"] for it in (fresh if i > 1 else page_items) if it["date"]]
        # 이 페이지 글이 전부 기준일보다 오래됐으면 그만
        if i >= min_pages and (not cutoff or not dates or max(dates) < cutoff):
            break
        # 이 페이지에 새로 나온 글을 이미 다 모아 뒀으면 그 뒤도 모아 둔 것이다 → 그만 (매번 8개월치를 다시 읽지 않게)
        check = fresh if i > 1 else page_items
        if i >= min_pages and is_known and check and all(
                is_known(it["url"] if it["url"] != url else it["title"]) for it in check):
            break
    return found


# ================================================================ watch
def page_lines(html: str, src: dict) -> list[str]:
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "noscript", "header", "footer", "nav", "form", "svg"]):
        tag.decompose()
    for sel in src.get("remove_selectors") or []:
        for tag in soup.select(sel):
            tag.decompose()
    root = None
    for sel in [src.get("content_selector"), "main", "article", "#content", "body"]:
        if sel and (root := soup.select_one(sel)) is not None:
            break
    if root is None:
        return []
    ignore = re.compile(src["ignore_regex"]) if src.get("ignore_regex") else None
    lines, seen = [], set()
    for raw in root.get_text("\n").split("\n"):
        line = clean(raw)
        if len(line) < 2 or line in seen or (ignore and ignore.search(line)):
            continue
        seen.add(line)
        lines.append(line)
    return lines


def snap_path(name: str) -> Path:
    return SNAP_DIR / (re.sub(r"[^\w가-힣ぁ-んァ-ン一-龥.-]+", "_", name) + ".txt")


def crawl_watch(src: dict, session: requests.Session) -> list[dict]:
    url = src["url"]
    if not robots_allows(url, session):
        raise PermissionError(f"robots.txt 가 수집을 막고 있다: {url}")
    lines = page_lines(fetch(url, src, session), src)
    if not lines:
        raise ValueError("본문을 찾지 못했다. content_selector 를 확인할 것")
    SNAP_DIR.mkdir(exist_ok=True)
    path = snap_path(src["name"])
    if not path.exists():  # 첫 실행: 기준 사본만 저장
        path.write_text("\n".join(lines), encoding="utf-8")
        return []
    old_set = set(path.read_text(encoding="utf-8").split("\n"))
    added = [l for l in lines if l not in old_set]
    removed_n = len(old_set - set(lines))
    path.write_text("\n".join(lines), encoding="utf-8")
    if not added and removed_n == 0:
        return []
    kws = src.get("keywords") or []
    if kws:
        added = [l for l in added if any(k.lower() in l.lower() for k in kws)]
        if not added:
            return []
    summary = "\n".join(added[:12]) if added else f"(문장 {removed_n}개 삭제)"
    if len(added) > 12:
        summary += f"\n…외 {len(added) - 12}줄"
    digest = hashlib.md5("\n".join(lines).encode()).hexdigest()[:10]
    return [{"title": f"{src['name']} 페이지가 바뀌었다",
             "date": datetime.now().strftime("%Y-%m-%d"), "url": url,
             "key": f"W|{url}|{digest}", "tkey": None, "change": summary, "kind": "watch"}]


# ================================================================ 미리보기 / 자동 추정 (저장 전 확인용)
def new_session() -> requests.Session:
    s = requests.Session()
    s.headers["User-Agent"] = USER_AGENT
    return s


def test_source(raw: dict) -> dict:
    """저장하지 않고 한 번 읽어 본다. 기록·사본은 건드리지 않는다."""
    src = clean_source(raw)
    session = new_session()
    _robots_cache.clear()
    url = page_url(src, 1) if src["type"] == "board" else src["url"]
    if not robots_allows(url, session):
        return {"ok": False, "error": "이 사이트의 robots.txt 가 수집을 금지한다. 추가하지 않는 것이 맞다."}
    try:
        html = fetch(url, src, session)
    except Exception as e:
        return {"ok": False, "error": f"페이지를 열지 못했다: {type(e).__name__}: {e}"[:300]}
    if src["type"] == "board":
        items = [it for it in parse_list(html, url, src) if keyword_ok(it["title"], src)]
        warn = []
        if not items:
            warn.append("글을 하나도 찾지 못했다. 선택자를 바꾸거나 '구조 자동 추정'을 눌러 볼 것.")
        elif not any(it["date"] for it in items):
            warn.append("날짜를 찾지 못했다. 목록은 쓸 수 있지만 날짜 칸이 비게 된다.")
        if items and all(it["url"] == url for it in items):
            warn.append("글 링크가 자바스크립트라 상세 주소를 만들지 못했다. 링크 규칙(link_regex/link_template)이 필요하다.")
        return {"ok": bool(items), "kind": "board", "count": len(items), "items": items[:15], "warnings": warn,
                "js_rendered": "<div id=\"app\"></div>" in html.replace(" ", "")}
    lines = page_lines(html, src)
    return {"ok": bool(lines), "kind": "watch", "count": len(lines), "lines": lines[:25],
            "warnings": [] if lines else ["본문을 찾지 못했다. 본문 선택자를 비우거나 바꿔 볼 것."]}


def _css_of(el) -> str:
    if el.get("id") and re.match(r"^[A-Za-z][\w-]*$", el["id"]):
        return f"{el.name}#{el['id']}"
    classes = [c for c in el.get("class", []) if re.match(r"^[A-Za-z_][\w-]*$", c)][:2]
    return el.name + "".join("." + c for c in classes)


def guess_board(url: str) -> dict:
    """게시판 목록 페이지를 읽어 '글 한 줄'과 '제목 링크' 선택자를 추정한다."""
    session = new_session()
    _robots_cache.clear()
    if not robots_allows(url, session):
        return {"ok": False, "error": "이 사이트의 robots.txt 가 수집을 금지한다."}
    try:
        html = fetch_static(url.replace("{page}", "1").replace("{offset}", "0"), session)
    except Exception as e:
        return {"ok": False, "error": f"페이지를 열지 못했다: {type(e).__name__}: {e}"[:300]}
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "noscript", "header", "footer", "nav"]):
        tag.decompose()
    groups: dict[str, list] = {}
    for el in soup.find_all(["tr", "li", "article", "dl"]):
        a = max(el.find_all("a"), key=lambda x: len(clean(x.get_text())), default=None)
        if a is None or len(clean(a.get_text())) < 6 or not DATE_RE.search(el.get_text(" ")):
            continue
        if el.find(["tr", "li", "article"]):  # 안쪽에 또 목록이 있으면 바깥 껍데기
            continue
        parent = el.parent
        if parent.name == "tbody":
            table = parent.find_parent("table")
            key = f"{_css_of(table)} tbody > {el.name}" if table else f"tbody > {el.name}"
        else:
            key = f"{_css_of(parent)} > {el.name}"
        groups.setdefault(key, []).append((el, a))
    if not groups:
        return {"ok": False, "error": "날짜가 붙은 글 목록을 찾지 못했다. JavaScript 로 목록을 그리는 사이트일 수 있다. 그럴 땐 '페이지 변경 감지' 방식으로 추가할 것."}
    row_sel, rows = max(groups.items(), key=lambda kv: len(kv[1]))
    # 제목 링크: 줄마다 가장 긴 링크의 class 가 공통이면 그것을 쓴다
    cls_sets = [set(a.get("class", [])) for _, a in rows]
    common = set.intersection(*cls_sets) if cls_sets else set()
    title_sel = "a." + sorted(common)[0] if common else "a"
    guess = {"type": "board", "list_url": url, "row_selector": row_sel, "title_selector": title_sel}
    notes = [f"글 {len(rows)}줄을 찾았다."]
    hrefs = [a.get("href", "") for _, a in rows]
    if hrefs and sum(h.lower().startswith("javascript") for h in hrefs) > len(hrefs) / 2:
        m = re.search(r"view\('(\d+)'", hrefs[0])
        if m and "list.do" in url:
            base = re.sub(r"([?&])pageIndex=[^&]*&?", r"\1", url.replace("list.do", "view.do")).rstrip("?&")
            guess["link_regex"] = r"view\('(\d+)'"
            guess["link_template"] = base + ("&" if "?" in base else "?") + "boardId={0}"
            notes.append("글 링크가 자바스크립트라 상세 주소 규칙을 추정해 넣었다. 미리보기에서 링크가 열리는지 확인할 것.")
        else:
            notes.append("글 링크가 자바스크립트라 상세 주소를 만들 수 없다. 목록 페이지 주소로 연결된다.")
    return {"ok": True, "guess": guess, "notes": notes}


# ================================================================ 실행
def run_crawl(store: Store, cfg: dict, group: str | None = None,
              progress: Callable[[dict], None] | None = None,
              stop: threading.Event | None = None) -> int:
    """설정된 모든 대상을 확인하고 새 항목 수를 돌려준다."""
    delay = float(cfg.get("delay_seconds", 2))
    session = requests.Session()
    session.headers["User-Agent"] = USER_AGENT
    _robots_cache.clear()

    cutoff = cutoff_date(cfg)
    max_pages = int(cfg.get("max_pages", 25))
    recent = (datetime.now() - timedelta(days=14)).strftime("%Y-%m-%d")
    targets = [s for s in cfg["sources"]
               if s.get("enabled", True) is not False and (not group or s.get("group") == group)]
    total_new = 0
    for i, src in enumerate(targets, 1):
        if stop and stop.is_set():
            break
        if progress:
            progress({"current": src["name"], "done": i - 1, "total": len(targets)})
        first_time = not store.source_seen_before(src["name"])
        try:
            if src.get("type") == "watch":
                raw = crawl_watch(src, session)
                time.sleep(delay)
            else:
                raw = crawl_board(src, session, delay, cutoff, max_pages,
                                  None if first_time else store.has_key)
        except Exception as e:  # 한 사이트가 실패해도 나머지는 계속
            store.set_source(src, ok=False, error=f"{type(e).__name__}: {e}"[:300])
            continue
        new = 0
        for it in raw:
            it.update(group=src.get("group", ""), source=src["name"])
            # 처음 확인하는 대상의 기존 글, 그리고 범위를 넓혀 뒤늦게 채운 2주 넘은 글은 '읽음'으로 넣는다.
            old = bool(it.get("date")) and it["date"] < recent
            if store.add_item(it, mark_read=first_time or old):
                new += 0 if (first_time or old) else 1
        store.set_source(src, ok=True, new=new)
        total_new += new
    store.set_meta("last_run", now_str())
    if progress:
        progress({"current": "", "done": len(targets), "total": len(targets)})
    return total_new


def export_xlsx(store: Store, path: Path = XLSX_PATH) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "공지"
    header = ["발견일시", "분류", "출처", "등록일", "제목", "링크", "정리 / 변경 내용", "읽음", "별표"]
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for col, w in zip("ABCDEFGHI", [17, 12, 24, 12, 70, 45, 60, 6, 6]):
        ws.column_dimensions[col].width = w
    for it in store.items(limit=100000):
        pts = json.loads(it["summary"]) if it.get("summary") else []
        body = "\n".join("• " + p for p in pts) if pts else it["change"]
        title = f'{it["title_ko"]} ({it["title"]})' if it.get("title_ko") else it["title"]
        ws.append([it["found_at"], it["grp"], it["source"], it["date"], title, it["url"],
                   body, "O" if it["is_read"] else "", "★" if it["starred"] else ""])
        ws.cell(row=ws.max_row, column=6).hyperlink = it["url"]
        ws.cell(row=ws.max_row, column=6).style = "Hyperlink"
        ws.cell(row=ws.max_row, column=7).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{ws.max_row}"
    wb.save(path)
    return path


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--group", help="이 분류만 확인 (예: 일본)")
    ap.add_argument("--export", action="store_true", help="notices.xlsx 로 내보내기")
    args = ap.parse_args()
    store = Store()
    if args.export:
        print(f"내보냄 → {export_xlsx(store)}")
        return 0
    cfg = load_config()
    n = run_crawl(store, cfg, args.group,
                  progress=lambda p: p["current"] and print(f"  확인 중 {p['done'] + 1}/{p['total']}  {p['current']}"))
    print(f"\n새 항목 {n}건")
    import summarize
    st = summarize.summarize_pending(store, cfg, progress=lambda m: print("  " + m))
    print(f"공지 정리 {st['done']}건 ({'Claude 정리·번역' if st['mode'] == 'llm' else '본문 발췌'})")
    for it in store.items(unread=True, limit=50):
        print(f"  [{it['grp']}] {it['date']}  {it['title']}\n      {it['url']}")
    for s in store.sources():
        if s["last_error"]:
            print(f"  [오류] {s['name']}: {s['last_error']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
